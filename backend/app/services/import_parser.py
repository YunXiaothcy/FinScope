import csv
import io
import posixpath
from dataclasses import dataclass
from datetime import date, datetime
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import fromstring
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.exceptions import InvalidFileException

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
MAX_COLUMNS = 40
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


@dataclass
class ParsedFile:
    headers: list[str]
    rows: list[tuple[int, list[str], list[str]]]
    encoding: str = ""
    sheet_name: str = ""


def invalid(message: str):
    raise HTTPException(422, message)


def validate_filename(filename: str) -> str:
    if (
        not filename
        or len(filename) > 160
        or any(char in filename for char in "/\\\x00")
        or any(ord(char) < 32 for char in filename)
    ):
        invalid("文件名无效，请使用不含路径的 CSV/XLSX 文件名")
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix not in {"csv", "xlsx"}:
        invalid("只支持 .csv 和 .xlsx，不支持旧版 XLS、宏文件或图片")
    return suffix


def tabular(rows, encoding="", sheet_name="") -> ParsedFile:
    rows = list(rows)
    if not rows:
        invalid("文件没有表头和数据")
    headers = [cell.strip() for cell in rows[0][1]]
    if (
        not headers
        or len(headers) > MAX_COLUMNS
        or any(not value or len(value) > 80 for value in headers)
    ):
        invalid("首行须为非空表头，最多 40 列，每列表头最多 80 字")
    if len(headers) != len(set(headers)) or rows[0][2]:
        invalid("表头不能重复、包含公式或错误值")
    result = []
    for number, cells, errors in rows[1:]:
        if not any(cells) and not errors:
            continue
        if any(len(cell) > 1000 for cell in cells):
            invalid("单元格内容过长，最多 1000 字")
        if len(cells) != len(headers):
            errors = [*errors, "该行列数与表头不一致"]
        result.append((number, cells, errors))
    if not result:
        invalid("文件没有可预览的数据行")
    return ParsedFile(headers, result, encoding, sheet_name)


def parse_csv(data: bytes) -> ParsedFile:
    text = None
    encoding = ""
    for candidate in ("utf-8-sig", "gb18030"):
        try:
            text = data.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if text is None or any(ord(char) < 32 and char not in "\t\r\n" for char in text):
        invalid("CSV 必须是 UTF-8 或 GB18030 文本，不能包含二进制控制字符")
    try:
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = []
        reader = csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True)
        for number, values in enumerate(reader, 1):
            if number > MAX_ROWS + 1 or len(values) > MAX_COLUMNS:
                invalid("每批最多 1000 行数据和 40 列，请拆分文件")
            rows.append((number, values, []))
        return tabular(rows, encoding=encoding)
    except csv.Error as exc:
        raise HTTPException(422, "CSV 格式无效，请检查分隔符和引号") from exc


def parse_xlsx(data: bytes) -> ParsedFile:
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 1000
                or sum(entry.file_size for entry in entries) > 10 * 1024 * 1024
            ):
                invalid("XLSX 解压后过大或文件项过多")
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)) or any(
                entry.flag_bits & 1 for entry in entries
            ):
                invalid("不支持加密或结构异常的 XLSX")
            if any(
                "vbaproject" in name.lower()
                or "externallinks/" in name.lower()
                or "embeddings/" in name.lower()
                for name in names
            ):
                invalid("请移除宏、外部链接和嵌入对象后再上传")
            # 所有 XML 在交给工作簿解析器前先拒绝实体、DTD 等不可信结构。
            for name in names:
                if name.endswith((".xml", ".rels")):
                    fromstring(archive.read(name), forbid_dtd=True)
            book_xml = fromstring(archive.read("xl/workbook.xml"))
            sheet = book_xml.find(f"{NS}sheets/{NS}sheet")
            if sheet is None:
                invalid("XLSX 没有工作表")
            if not sheet.get("name") or len(sheet.get("name")) > 80:
                invalid("工作表名称无效或过长，请在来源文件中修改名称")
            rel_id = sheet.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            )
            rels = fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            target = next(
                (
                    rel.get("Target", "")
                    for rel in rels
                    if rel.get("Id") == rel_id and rel.get("TargetMode") != "External"
                ),
                "",
            )
            path = posixpath.normpath(
                target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            )
            if not path.startswith("xl/worksheets/") or path not in names:
                invalid("第一张表必须是普通数据工作表")
            worksheet_xml = fromstring(archive.read(path))
            numeric = {}
            max_row, max_column = 1, 1
            for cell in worksheet_xml.iter(f"{NS}c"):
                row, column = coordinate_to_tuple(cell.get("r", ""))
                max_row, max_column = max(max_row, row), max(max_column, column)
                if max_row > MAX_ROWS + 1 or max_column > MAX_COLUMNS:
                    invalid(
                        "第一张工作表最多 1000 行数据和 40 列，请移除多余内容或拆分"
                    )
                value = cell.find(f"{NS}v")
                if cell.get("t", "n") == "n" and value is not None:
                    numeric[cell.get("r")] = value.text or ""
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            worksheet = workbook.worksheets[0]
            rows = []
            for number, cells in enumerate(
                worksheet.iter_rows(max_row=max_row, max_col=max_column), 1
            ):
                values, errors = [], []
                for cell in cells:
                    value = cell.value
                    if cell.data_type in {"f", "e"}:
                        errors.append(
                            "该行含公式或 Excel 错误值，请在原表中粘贴为确定的值"
                        )
                    if isinstance(value, datetime | date):
                        text = value.isoformat()
                    elif getattr(cell, "coordinate", None) in numeric:
                        # OOXML 十进制原文直接用于金额，避免经过 openpyxl 的 float 值。
                        text = numeric[cell.coordinate]
                    else:
                        text = "" if value is None else str(value)
                    values.append(text)
                rows.append((number, values, list(dict.fromkeys(errors))))
            return tabular(rows, sheet_name=worksheet.title)
        finally:
            workbook.close()
    except HTTPException:
        raise
    except (
        BadZipFile,
        KeyError,
        ValueError,
        TypeError,
        OverflowError,
        OSError,
        DefusedXmlException,
        ParseError,
        InvalidFileException,
        IndexError,
        NotImplementedError,
    ) as exc:
        raise HTTPException(
            422, "XLSX 内容无效或包含不安全结构，请重新另存为普通 XLSX"
        ) from exc


def parse_file(filename: str, data: bytes) -> ParsedFile:
    suffix = validate_filename(filename)
    if not data or len(data) > MAX_BYTES:
        raise HTTPException(413, "文件不能为空且不得超过 2 MiB")
    return parse_csv(data) if suffix == "csv" else parse_xlsx(data)
